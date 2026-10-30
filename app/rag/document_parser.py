"""
Document parsing module.
Supports PDF (PyMuPDF), DOCX (python-docx), TXT, and Markdown files.
"""
from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class ParsedChunk:
    """A single text chunk with source metadata."""
    text: str
    source_file: str
    page: Optional[int] = None
    section: Optional[str] = None
    chunk_index: int = 0

    def to_dict(self) -> dict:
        return {
            "text": self.text,
            "source_file": self.source_file,
            "page": self.page,
            "section": self.section,
            "chunk_index": self.chunk_index,
        }


@dataclass
class ParseResult:
    """Result of parsing a document."""
    full_text: str = ""
    chunks: list[ParsedChunk] = field(default_factory=list)
    technologies: list[str] = field(default_factory=list)
    error: Optional[str] = None
    warning: Optional[str] = None
    page_count: int = 0


def _detect_technologies(text: str) -> list[str]:
    """Heuristic detection of technologies mentioned in document text."""
    patterns = [
        # Languages
        r"\bPython\b", r"\bJavaScript\b", r"\bTypeScript\b", r"\bJava\b",
        r"\bC\+\+\b", r"\bC#\b", r"\bRust\b", r"\bGo\b", r"\bKotlin\b",
        r"\bSwift\b", r"\bRuby\b", r"\bPHP\b", r"\bScala\b", r"\bR\b",
        # Frameworks / Libraries
        r"\bFlask\b", r"\bDjango\b", r"\bFastAPI\b", r"\bReact\b",
        r"\bVue\b", r"\bAngular\b", r"\bNext\.js\b", r"\bNode\.js\b",
        r"\bSpring\b", r"\bTensorFlow\b", r"\bPyTorch\b", r"\bScikit-learn\b",
        r"\bPandas\b", r"\bNumPy\b", r"\bOpenCV\b", r"\bHugging Face\b",
        r"\bLangChain\b", r"\bFAISS\b", r"\bSentence[- ]Transformers\b",
        # Databases
        r"\bPostgreSQL\b", r"\bMySQL\b", r"\bSQLite\b", r"\bMongoDB\b",
        r"\bRedis\b", r"\bElasticsearch\b", r"\bCassandra\b", r"\bSQLAlchemy\b",
        # Cloud / Infra
        r"\bAWS\b", r"\bGCP\b", r"\bAzure\b", r"\bDocker\b", r"\bKubernetes\b",
        r"\bTerraform\b", r"\bGitHub Actions\b", r"\bCI/CD\b",
        # Concepts
        r"\bRAG\b", r"\bLLM\b", r"\bGPT\b", r"\bBERT\b", r"\bTransformer\b",
        r"\bREST\b", r"\bGraphQL\b", r"\bgRPC\b", r"\bWebSocket\b",
        r"\bMicroservice\b", r"\bMessage Queue\b", r"\bKafka\b", r"\bRabbitMQ\b",
    ]
    found = set()
    for pat in patterns:
        if re.search(pat, text, re.IGNORECASE):
            # Clean up the matched word
            match = re.search(pat, text, re.IGNORECASE)
            if match:
                found.add(match.group(0).strip())
    return sorted(found)


def _chunk_text(
    text: str,
    source_file: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
    page: Optional[int] = None,
    section: Optional[str] = None,
    start_index: int = 0,
) -> list[ParsedChunk]:
    """Split text into overlapping chunks."""
    chunks = []
    words = text.split()
    if not words:
        return chunks

    step = max(1, chunk_size - chunk_overlap)
    idx = 0
    chunk_num = start_index

    while idx < len(words):
        chunk_words = words[idx: idx + chunk_size]
        chunk_text = " ".join(chunk_words).strip()
        if chunk_text:
            chunks.append(ParsedChunk(
                text=chunk_text,
                source_file=source_file,
                page=page,
                section=section,
                chunk_index=chunk_num,
            ))
            chunk_num += 1
        idx += step

    return chunks


def parse_pdf(file_path: str, chunk_size: int = 800, chunk_overlap: int = 150) -> ParseResult:
    """Parse a PDF using PyMuPDF, page by page."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ParseResult(error="PyMuPDF not installed.")

    result = ParseResult()
    filename = Path(file_path).name
    full_text_parts: list[str] = []
    chunk_idx = 0

    try:
        doc = fitz.open(file_path)
        result.page_count = doc.page_count

        # Check if the PDF appears to be scanned (text-poor)
        total_text_len = sum(len(page.get_text("text")) for page in doc)
        if total_text_len < 200 and doc.page_count > 0:
            result.warning = (
                "This PDF appears to contain scanned images with little extractable text. "
                "OCR is required for accurate processing. Consider converting the PDF with "
                "an OCR tool (e.g., Adobe Acrobat, Tesseract) before uploading."
            )

        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text("text")
            if not page_text.strip():
                continue
            full_text_parts.append(page_text)
            chunks = _chunk_text(
                page_text, filename,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                page=page_num,
                start_index=chunk_idx,
            )
            result.chunks.extend(chunks)
            chunk_idx += len(chunks)

        doc.close()
        result.full_text = "\n".join(full_text_parts)
        result.technologies = _detect_technologies(result.full_text)

    except Exception as exc:
        logger.exception("PDF parsing failed: %s", exc)
        result.error = f"Failed to parse PDF: {exc}"

    return result


def parse_docx(file_path: str, chunk_size: int = 800, chunk_overlap: int = 150) -> ParseResult:
    """Parse a DOCX file using python-docx."""
    try:
        from docx import Document
    except ImportError:
        return ParseResult(error="python-docx not installed.")

    result = ParseResult()
    filename = Path(file_path).name
    full_text_parts: list[str] = []
    current_section = "Document"
    section_buffer: list[str] = []
    chunk_idx = 0

    try:
        doc = Document(file_path)

        def flush_section():
            nonlocal chunk_idx
            text = "\n".join(section_buffer).strip()
            if text:
                chunks = _chunk_text(
                    text, filename,
                    chunk_size=chunk_size,
                    chunk_overlap=chunk_overlap,
                    section=current_section,
                    start_index=chunk_idx,
                )
                result.chunks.extend(chunks)
                chunk_idx += len(chunks)
                section_buffer.clear()

        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            full_text_parts.append(text)
            if style.startswith("Heading"):
                flush_section()
                current_section = text  # type: ignore[assignment]
                section_buffer.append(text)
            else:
                section_buffer.append(text)

        flush_section()

        # Also extract tables
        for table in doc.tables:
            rows = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    rows.append(row_text)
            table_text = "\n".join(rows)
            if table_text:
                full_text_parts.append(table_text)
                chunks = _chunk_text(
                    table_text, filename,
                    chunk_size=chunk_size,
                    chunk_overlap=chunk_overlap,
                    section="Table",
                    start_index=chunk_idx,
                )
                result.chunks.extend(chunks)
                chunk_idx += len(chunks)

        result.full_text = "\n".join(full_text_parts)
        result.technologies = _detect_technologies(result.full_text)

    except Exception as exc:
        logger.exception("DOCX parsing failed: %s", exc)
        result.error = f"Failed to parse DOCX: {exc}"

    return result


def parse_text(file_path: str, chunk_size: int = 800, chunk_overlap: int = 150) -> ParseResult:
    """Parse plain text or Markdown files."""
    result = ParseResult()
    filename = Path(file_path).name

    try:
        with open(file_path, encoding="utf-8", errors="replace") as f:
            text = f.read()

        result.full_text = text
        result.technologies = _detect_technologies(text)

        # For Markdown, try to detect sections by headings
        sections: list[tuple[str, str]] = []
        current_heading = "Document"
        buffer: list[str] = []

        for line in text.splitlines():
            heading_match = re.match(r"^(#{1,6})\s+(.*)", line)
            if heading_match:
                if buffer:
                    sections.append((current_heading, "\n".join(buffer)))
                    buffer = []
                current_heading = heading_match.group(2).strip()
                buffer.append(line)
            else:
                buffer.append(line)

        if buffer:
            sections.append((current_heading, "\n".join(buffer)))

        chunk_idx = 0
        for section_name, section_text in sections:
            chunks = _chunk_text(
                section_text, filename,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                section=section_name,
                start_index=chunk_idx,
            )
            result.chunks.extend(chunks)
            chunk_idx += len(chunks)

    except Exception as exc:
        logger.exception("Text parsing failed: %s", exc)
        result.error = f"Failed to parse text file: {exc}"

    return result


def parse_document(
    file_path: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
) -> ParseResult:
    """Route to the correct parser based on file extension."""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path, chunk_size, chunk_overlap)
    elif ext == ".docx":
        return parse_docx(file_path, chunk_size, chunk_overlap)
    elif ext in (".txt", ".md", ".markdown"):
        return parse_text(file_path, chunk_size, chunk_overlap)
    else:
        return ParseResult(error=f"Unsupported file type: {ext}")
