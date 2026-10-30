"""
Automated tests for AI Viva Simulator.
Run: pytest tests/ -v
"""
import io
import json
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

# Ensure project root is on path
sys.path.insert(0, str(Path(__file__).parent.parent))

os.environ.setdefault("GROQ_API_KEY", "test-key-not-real")
os.environ.setdefault("DATABASE_URL", "sqlite:///test_viva.db")
os.environ.setdefault("UPLOAD_FOLDER", "test_uploads")
os.environ.setdefault("FAISS_INDEX_DIR", "test_faiss_indexes")


# ── App fixture ────────────────────────────────────────────
@pytest.fixture(scope="module")
def app():
    from app_factory import create_app
    app = create_app()
    app.config["TESTING"] = True
    app.config["WTF_CSRF_ENABLED"] = False
    return app


@pytest.fixture(scope="module")
def client(app):
    return app.test_client()


@pytest.fixture(scope="module", autouse=True)
def cleanup():
    yield
    # Clean up test DB
    db_path = Path("test_viva.db")
    if db_path.exists():
        try:
            db_path.unlink()
        except Exception:
            pass
    for d in ["test_uploads", "test_faiss_indexes"]:
        import shutil
        if Path(d).exists():
            try:
                shutil.rmtree(d)
            except Exception:
                pass


# ── Test: Document parsing ─────────────────────────────────
class TestDocumentParser:

    def test_parse_txt(self, tmp_path):
        from app.rag.document_parser import parse_text
        txt_file = tmp_path / "sample.txt"
        txt_file.write_text(
            "This project uses Python and Flask for the backend.\n"
            "The database is SQLite managed via SQLAlchemy.\n"
            "We use FAISS for vector search and sentence-transformers for embeddings.\n" * 20
        )
        result = parse_text(str(txt_file), chunk_size=50, chunk_overlap=10)
        assert result.error is None, f"Parse error: {result.error}"
        assert len(result.full_text) > 0
        assert len(result.chunks) >= 1
        assert "Python" in result.technologies or "Flask" in result.technologies

    def test_parse_txt_tech_detection(self, tmp_path):
        from app.rag.document_parser import _detect_technologies
        text = "We used React for frontend, FastAPI for backend, and PostgreSQL for the database."
        techs = _detect_technologies(text)
        assert any("React" in t for t in techs)
        assert any("FastAPI" in t for t in techs)
        assert any("PostgreSQL" in t for t in techs)

    def test_parse_markdown(self, tmp_path):
        from app.rag.document_parser import parse_text
        md_file = tmp_path / "readme.md"
        content = """# My Project
## Architecture
This uses a microservice architecture with Docker containers.
## Technologies
- Python 3.11
- Flask web framework
- Redis for caching
""" * 5
        md_file.write_text(content)
        result = parse_text(str(md_file))
        assert result.error is None
        assert len(result.chunks) >= 1
        # Sections should be detected
        sections = [c.section for c in result.chunks if c.section]
        assert len(sections) >= 1

    def test_parse_docx(self, tmp_path):
        """Test DOCX parsing creates chunks with section metadata."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        from app.rag.document_parser import parse_docx

        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This project implements a student management system using Python.")
        doc.add_heading("Technology Stack", level=2)
        doc.add_paragraph("Flask is used as the web framework." * 30)
        doc_path = tmp_path / "test.docx"
        doc.save(str(doc_path))

        result = parse_docx(str(doc_path))
        assert result.error is None
        assert len(result.full_text) > 10
        assert len(result.chunks) >= 1

    def test_parse_unsupported(self, tmp_path):
        from app.rag.document_parser import parse_document
        result = parse_document("fake_file.xyz")
        assert result.error is not None
        assert "Unsupported" in result.error

    def test_chunk_overlap(self, tmp_path):
        from app.rag.document_parser import _chunk_text
        text = "word " * 200
        chunks = _chunk_text(text, "test.txt", chunk_size=50, chunk_overlap=10)
        assert len(chunks) >= 2
        # Verify overlap: last words of chunk N should appear at start of chunk N+1
        c1_words = set(chunks[0].text.split())
        c2_words = set(chunks[1].text.split())
        assert len(c1_words & c2_words) > 0  # overlap exists


# ── Test: Vector store isolation ──────────────────────────
class TestVectorStoreIsolation:

    def test_separate_indexes(self, tmp_path):
        """Verifies that two projects have fully isolated indexes."""
        from app.rag.document_parser import ParsedChunk
        from app.rag.vector_store import ProjectVectorStore

        with patch("app.rag.vector_store.Config") as mock_config:
            mock_config.FAISS_INDEX_DIR = str(tmp_path)
            mock_config.EMBEDDING_MODEL = "all-MiniLM-L6-v2"

            chunks_a = [ParsedChunk(text="Python Flask web application backend API", source_file="a.txt", chunk_index=0)]
            chunks_b = [ParsedChunk(text="React JavaScript frontend user interface", source_file="b.txt", chunk_index=0)]

            store_a = ProjectVectorStore(1001)
            store_a.index_dir = tmp_path / "project_1001"
            store_a.index_dir.mkdir(exist_ok=True)
            store_a.index_path = store_a.index_dir / "index.faiss"
            store_a.meta_path = store_a.index_dir / "chunks.pkl"

            store_b = ProjectVectorStore(1002)
            store_b.index_dir = tmp_path / "project_1002"
            store_b.index_dir.mkdir(exist_ok=True)
            store_b.index_path = store_b.index_dir / "index.faiss"
            store_b.meta_path = store_b.index_dir / "chunks.pkl"

            store_a.build(chunks_a)
            store_b.build(chunks_b)

            # Retrieve from A – should only find A's content
            results_a = store_a.retrieve("Python Flask", top_k=1)
            results_b = store_b.retrieve("React JavaScript", top_k=1)

            assert results_a, "Store A should return results"
            assert results_b, "Store B should return results"
            assert results_a[0]["source_file"] == "a.txt"
            assert results_b[0]["source_file"] == "b.txt"


# ── Test: API endpoints ────────────────────────────────────
class TestProjectAPI:

    def test_list_projects_empty(self, client):
        res = client.get("/api/projects")
        assert res.status_code == 200
        data = res.get_json()
        assert "projects" in data
        assert isinstance(data["projects"], list)

    def test_upload_no_file(self, client):
        res = client.post("/api/projects")
        assert res.status_code == 400
        data = res.get_json()
        assert "error" in data

    def test_upload_invalid_extension(self, client):
        data = {"file": (io.BytesIO(b"content"), "file.exe")}
        res = client.post("/api/projects", data=data, content_type="multipart/form-data")
        assert res.status_code == 400

    def test_upload_valid_txt(self, client):
        """Upload a valid TXT file and check processing starts."""
        content = b"This is a Flask project using Python and SQLite. " * 50
        data = {
            "file": (io.BytesIO(content), "test_project.txt"),
            "title": "Test Project",
        }
        res = client.post("/api/projects", data=data, content_type="multipart/form-data")
        assert res.status_code == 201
        resp = res.get_json()
        assert "project" in resp
        assert resp["project"]["title"] == "Test Project"
        assert resp["project"]["processing_status"] == "pending"

    def test_get_project_not_found(self, client):
        res = client.get("/api/projects/99999")
        assert res.status_code == 404

    def test_delete_project_not_found(self, client):
        res = client.delete("/api/projects/99999")
        assert res.status_code == 404


class TestSessionAPI:

    def test_create_session_no_project(self, client):
        res = client.post("/api/sessions",
                          data=json.dumps({"project_id": 99999}),
                          content_type="application/json")
        assert res.status_code == 404

    def test_create_session_missing_project_id(self, client):
        res = client.post("/api/sessions",
                          data=json.dumps({}),
                          content_type="application/json")
        assert res.status_code == 400

    def test_get_session_not_found(self, client):
        res = client.get("/api/sessions/99999")
        assert res.status_code == 404

    def test_submit_answer_invalid_session(self, client):
        res = client.post("/api/sessions/99999/answer",
                          data=json.dumps({"question_id": 1, "answer": "test"}),
                          content_type="application/json")
        assert res.status_code == 400


# ── Test: LLM JSON parsing ─────────────────────────────────
class TestLLMSafeParse:

    def test_clean_json(self):
        from app.llm.examiner import _safe_json_parse
        result = _safe_json_parse('{"score": 75, "correctness": "partial"}')
        assert result is not None
        assert result["score"] == 75

    def test_json_in_code_fence(self):
        from app.llm.examiner import _safe_json_parse
        raw = '```json\n{"score": 80, "correctness": "correct"}\n```'
        result = _safe_json_parse(raw)
        assert result is not None
        assert result["score"] == 80

    def test_json_with_prose(self):
        from app.llm.examiner import _safe_json_parse
        raw = 'Here is the evaluation:\n{"score": 60, "correctness": "partial"}\nThat is all.'
        result = _safe_json_parse(raw)
        assert result is not None
        assert result["score"] == 60

    def test_trailing_comma(self):
        from app.llm.examiner import _safe_json_parse
        raw = '{"score": 50, "correctness": "incorrect",}'
        result = _safe_json_parse(raw)
        assert result is not None

    def test_invalid_json_returns_none(self):
        from app.llm.examiner import _safe_json_parse
        result = _safe_json_parse("This is not JSON at all.")
        assert result is None


# ── Test: Speech service ───────────────────────────────────
class TestSpeechService:

    def test_transcribe_no_key(self):
        from app.speech.transcription import transcribe_audio
        with patch("app.speech.transcription.Config") as mock_cfg:
            mock_cfg.GROQ_API_KEY = ""
            result = transcribe_audio(b"fake audio bytes")
        assert result["transcript"] == ""
        assert "GROQ_API_KEY" in result["error"]

    def test_transcribe_empty_audio(self):
        from app.speech.transcription import transcribe_audio
        result = transcribe_audio(b"")
        assert result["transcript"] == ""
        assert result["error"] is not None

    def test_transcribe_too_large(self):
        from app.speech.transcription import transcribe_audio, MAX_AUDIO_BYTES
        huge = b"x" * (MAX_AUDIO_BYTES + 1)
        result = transcribe_audio(huge)
        assert "too large" in result["error"].lower()


# ── Test: Speech API endpoint ──────────────────────────────
class TestSpeechEndpoint:

    def test_transcribe_no_audio_field(self, client):
        res = client.post("/api/speech/transcribe")
        assert res.status_code == 400

    def test_transcribe_empty_file(self, client):
        data = {"audio": (io.BytesIO(b""), "recording.webm")}
        res = client.post("/api/speech/transcribe", data=data, content_type="multipart/form-data")
        assert res.status_code in (400, 503)


# ── Test: Frontend routes ──────────────────────────────────
class TestFrontendRoutes:

    def test_index(self, client):
        res = client.get("/")
        assert res.status_code == 200
        assert b"Viva Simulator" in res.data or b"viva" in res.data.lower()

    def test_viva_page(self, client):
        res = client.get("/viva/1")
        assert res.status_code == 200

    def test_report_page(self, client):
        res = client.get("/report/1")
        assert res.status_code == 200

    def test_404_api(self, client):
        res = client.get("/api/nonexistent")
        assert res.status_code == 404
        data = res.get_json()
        assert "error" in data
